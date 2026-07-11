# -*- coding: utf-8 -*-
"""
将 output/Agentic原生架构图.html（Architecture Diagram Generator 风格）转为一页可编辑 PPT。
深色 slate-950 主题：语义配色（青=交互 / 翠=LLM·Agent / 紫=知识 / 橙=总线 / 玫瑰=安全 / 琥珀=边界）。
左侧为 SVG 架构图的等比还原（全部原生形状），右侧为三张摘要卡片。
运行：python build_agentic_arch_ppt.py  →  output/Agentic原生架构图.pptx
"""
import os
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

from ppt_kit import new_prs, slide, txt, R, save

OUT = "/Users/apple/SocratesAI/output"
In = Inches

# ---------- 配色（HTML 同款；半透明填充按深色底预混） ----------
BG      = RGBColor(0x02, 0x06, 0x17)   # slate-950
PANEL   = RGBColor(0x0F, 0x17, 0x2A)   # slate-900
BORDER  = RGBColor(0x1E, 0x29, 0x3B)   # slate-800
WHITE   = RGBColor(0xFF, 0xFF, 0xFF)
SLATE   = RGBColor(0x94, 0xA3, 0xB8)   # 副标签
SLATE6  = RGBColor(0x64, 0x74, 0x8B)   # 中性箭头
FAINT   = RGBColor(0x47, 0x55, 0x69)
CYAN    = RGBColor(0x22, 0xD3, 0xEE)
EMERALD = RGBColor(0x34, 0xD3, 0x99)
VIOLET  = RGBColor(0xA7, 0x8B, 0xFA)
AMBER   = RGBColor(0xFB, 0xBF, 0x24)
ROSE    = RGBColor(0xFB, 0x71, 0x85)
ORANGE  = RGBColor(0xFB, 0x92, 0x3C)
F_CYAN    = RGBColor(0x0C, 0x22, 0x34)   # rgba(8,51,68,.4)   over #0f172a
F_EMERALD = RGBColor(0x0B, 0x2D, 0x31)   # rgba(6,78,59,.4)
F_VIOLET  = RGBColor(0x27, 0x19, 0x55)   # rgba(76,29,149,.4)
F_ROSE    = RGBColor(0x3F, 0x15, 0x2F)   # rgba(136,19,55,.4)
F_ORANGE  = RGBColor(0x56, 0x3C, 0x2F)   # rgba(251,146,60,.3)
F_GENERIC = RGBColor(0x17, 0x20, 0x33)   # rgba(30,41,59,.5)
F_HARNESS = RGBColor(0x0E, 0x0F, 0x18)   # rgba(251,191,36,.05) over #020617

# ---------- SVG(1000x668) → 幻灯片坐标 ----------
KX, KY = 9.2 / 1000.0, 6.05 / 668.0
X0, Y0 = 0.35, 1.02
def fx(px): return In(X0 + px * KX)
def fy(px): return In(Y0 + px * KY)
def fw(px): return In(px * KX)
def fh(px): return In(px * KY)


def _dash(shp, val="dash"):
    ln = shp.line._get_or_add_ln()
    ln.append(ln.makeelement(qn('a:prstDash'), {'val': val}))


def box(s, x, y, w, h, fill, line_c=None, line_w=1.2, dash=None, round_=True):
    shp = s.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if round_ else MSO_SHAPE.RECTANGLE, x, y, w, h)
    if round_:
        try:
            shp.adjustments[0] = 0.10
        except Exception:
            pass
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid(); shp.fill.fore_color.rgb = fill
    if line_c is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line_c; shp.line.width = Pt(line_w)
        if dash:
            _dash(shp, dash)
    shp.shadow.inherit = False
    return shp


def arrow(s, shape, x, y, w, h, fill):
    a = s.shapes.add_shape(shape, x, y, w, h)
    a.fill.solid(); a.fill.fore_color.rgb = fill
    a.line.fill.background()
    a.shadow.inherit = False
    return a


def node(s, px, py, pw, ph, name, sub, fill, stroke, name_sz=9, sub_sz=7.5):
    box(s, fx(px), fy(py), fw(pw), fh(ph), fill, stroke)
    paras = [[R(name, name_sz, WHITE, True)]]
    if sub:
        paras.append([R(sub, sub_sz, SLATE)])
    txt(s, fx(px), fy(py), fw(pw), fh(ph), paras,
        align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, space_after=0, line_spacing=1.05)


def card(s, x, y, w, h, dot, title, items):
    box(s, x, y, w, h, PANEL, BORDER, 1.0)
    d = s.shapes.add_shape(MSO_SHAPE.OVAL, x + In(0.16), y + In(0.16), In(0.10), In(0.10))
    d.fill.solid(); d.fill.fore_color.rgb = dot
    d.line.fill.background(); d.shadow.inherit = False
    txt(s, x + In(0.34), y + In(0.09), w - In(0.45), In(0.26),
        [[R(title, 10, WHITE, True)]], anchor=MSO_ANCHOR.MIDDLE)
    txt(s, x + In(0.18), y + In(0.42), w - In(0.34), h - In(0.52),
        [[R("• " + it, 8, SLATE)] for it in items], line_spacing=1.12, space_after=4)


def build():
    prs = new_prs()
    s = slide(prs)
    SW, SH = In(13.333), In(7.5)
    box(s, 0, 0, SW, SH, BG, round_=False)

    # ---------- 标题 ----------
    d = s.shapes.add_shape(MSO_SHAPE.OVAL, In(0.38), In(0.30), In(0.13), In(0.13))
    d.fill.solid(); d.fill.fore_color.rgb = CYAN; d.line.fill.background(); d.shadow.inherit = False
    txt(s, In(0.62), In(0.16), In(9.5), In(0.36), [[R("Agentic AI · 原生架构图", 16, WHITE, True)]])
    txt(s, In(0.62), In(0.52), In(12.3), In(0.28),
        [[R("场景 1 LLM+MCP 工具 → 场景 2 LLM+RAG/SKILL → 场景 3 多 Agent 协同 —— 全部包含在场景 4 · Harness 之中",
            9, SLATE)]])

    # ---------- 场景 4：Harness 大框 ----------
    box(s, fx(160), fy(36), fw(815), fh(564), F_HARNESS, AMBER, 1.0, dash="dash")
    txt(s, fx(174), fy(42), fw(700), In(0.24),
        [[R("场景 4 · HARNESS —— Agent 运行框架：把场景 1 / 2 / 3 全部包含的大框", 10, AMBER, True)]])
    txt(s, fx(340), fy(84), fw(580), In(0.42),
        [[R("圈内的一切 —— 工具调用 · 知识与流程 · 多 Agent 协同 ——", 8, SLATE)],
         [R("都由同一运行框架统一承载，这就是 Harness 的内涵", 8, SLATE)]],
        align=PP_ALIGN.CENTER, space_after=2, line_spacing=1.1)

    # ---------- 场景边界 ----------
    for (bx, bw, c, lab, cap) in [
            (180, 245, CYAN,    "场景 1 · LLM + MCP 工具",   "简单问答 · 内容生成"),
            (445, 245, VIOLET,  "场景 2 · LLM + RAG / SKILL", "知识定边界 · SKILL 定章法"),
            (710, 240, EMERALD, "场景 3 · 多 Agent 协同",     "调度 · 协同 · 自适应")]:
        box(s, fx(bx), fy(170), fw(bw), fh(300), None, c, 1.0, dash="dash")
        txt(s, fx(bx + 10), fy(176), fw(bw - 20), In(0.22), [[R(lab, 8.5, c, True)]])
        txt(s, fx(bx), fy(440), fw(bw), In(0.20),
            [[R(cap, 7.5, c)]], align=PP_ALIGN.CENTER)

    # ---------- 外部用户与入口 ----------
    node(s, 25, 82, 110, 52, "企业用户", "业务人员 / 系统", F_GENERIC, SLATE)
    node(s, 180, 82, 110, 52, "Agent 入口", "会话 · API · 任务", F_CYAN, CYAN)
    arrow(s, MSO_SHAPE.RIGHT_ARROW, fx(138), fy(93), fw(38), fh(14), CYAN)
    txt(s, fx(130), fy(72), fw(56), In(0.16), [[R("请求", 6.5, SLATE)]], align=PP_ALIGN.CENTER)
    a = arrow(s, MSO_SHAPE.LEFT_ARROW, fx(138), fy(115), fw(38), fh(14), SLATE6)
    txt(s, fx(130), fy(132), fw(56), In(0.16), [[R("结果", 6.5, SLATE)]], align=PP_ALIGN.CENTER)
    arrow(s, MSO_SHAPE.DOWN_ARROW, fx(228), fy(138), fw(14), fh(30), CYAN)

    # ---------- 场景 1：LLM ⇅ MCP 总线 → 工具 ----------
    node(s, 240, 200, 125, 54, "LLM", "理解 · 推理 · 生成", F_EMERALD, EMERALD)
    arrow(s, MSO_SHAPE.DOWN_ARROW, fx(289), fy(256), fw(13), fh(28), EMERALD)
    arrow(s, MSO_SHAPE.UP_ARROW,   fx(305), fy(256), fw(13), fh(28), ORANGE)
    box(s, fx(225), fy(286), fw(155), fh(22), F_ORANGE, ORANGE, 1.0)
    txt(s, fx(225), fy(286), fw(155), fh(22),
        [[R("MCP 协议 / 工具调用总线", 6.5, ORANGE, True)]],
        align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    for ax in (220, 296, 372):
        arrow(s, MSO_SHAPE.DOWN_ARROW, fx(ax), fy(310), fw(13), fh(28), ORANGE)
    node(s, 192, 340, 70, 52, "搜索", "Web 检索", F_GENERIC, SLATE, name_sz=8.5, sub_sz=7)
    node(s, 268, 340, 70, 52, "文件", "文档系统", F_GENERIC, SLATE, name_sz=8.5, sub_sz=7)
    node(s, 344, 340, 70, 52, "业务 API", "ERP · CRM", F_GENERIC, SLATE, name_sz=8.5, sub_sz=7)

    # ---------- 场景 2：知识库 / SKILL → LLM ----------
    node(s, 505, 200, 125, 54, "LLM", "按知识与流程作答", F_EMERALD, EMERALD)
    arrow(s, MSO_SHAPE.UP_ARROW, fx(502), fy(258), fw(13), fh(80), VIOLET)
    txt(s, fx(516), fy(288), fw(70), In(0.16), [[R("检索增强", 6.5, SLATE)]])
    arrow(s, MSO_SHAPE.UP_ARROW, fx(617), fy(258), fw(13), fh(80), CYAN)
    txt(s, fx(631), fy(288), fw(70), In(0.16), [[R("流程规范", 6.5, SLATE)]])
    node(s, 457, 340, 105, 56, "知识库", "RAG · 向量检索", F_VIOLET, VIOLET)
    node(s, 572, 340, 105, 56, "SKILL", "业务流程 · 规范", F_CYAN, CYAN)

    # ---------- 场景 3：调度 → 协同总线 → 执行 Agents + 自适应 ----------
    node(s, 762, 200, 136, 50, "调度 Agent", "规划 · 分派 · 汇聚", F_EMERALD, EMERALD)
    arrow(s, MSO_SHAPE.DOWN_ARROW, fx(823), fy(252), fw(13), fh(28), EMERALD)
    box(s, fx(735), fy(282), fw(190), fh(20), F_ORANGE, ORANGE, 1.0)
    txt(s, fx(735), fy(282), fw(190), fh(20),
        [[R("Agent 协同总线（A2A）", 6.5, ORANGE, True)]],
        align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    for ax in (741, 813, 885):
        arrow(s, MSO_SHAPE.DOWN_ARROW, fx(ax), fy(304), fw(13), fh(28), ORANGE)
    node(s, 718, 334, 60, 52, "检索", "Agent", F_EMERALD, EMERALD, name_sz=8.5, sub_sz=7)
    node(s, 790, 334, 60, 52, "分析", "Agent", F_EMERALD, EMERALD, name_sz=8.5, sub_sz=7)
    node(s, 862, 334, 60, 52, "执行", "Agent", F_EMERALD, EMERALD, name_sz=8.5, sub_sz=7)
    txt(s, fx(772), fy(350), fw(24), In(0.18), [[R("⇄", 8, EMERALD, True)]], align=PP_ALIGN.CENTER)
    txt(s, fx(844), fy(350), fw(24), In(0.18), [[R("⇄", 8, EMERALD, True)]], align=PP_ALIGN.CENTER)
    arrow(s, MSO_SHAPE.UP_ARROW, fx(930), fy(238), fw(12), fh(118), AMBER)
    txt(s, fx(945), In(fy(296).inches - 0.44), In(0.20), In(0.92),
        [[R(ch, 6.5, AMBER, True)] for ch in "自适应反馈"],
        align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, space_after=1, line_spacing=1.0)

    # ---------- 场景演进箭头 ----------
    arrow(s, MSO_SHAPE.RIGHT_ARROW, fx(427), fy(312), fw(17), fh(16), SLATE6)
    arrow(s, MSO_SHAPE.RIGHT_ARROW, fx(692), fy(312), fw(17), fh(16), SLATE6)

    # ---------- Harness 统一能力底座 ----------
    for ax in (295, 560, 823):
        arrow(s, MSO_SHAPE.DOWN_ARROW, fx(ax), fy(472), fw(13), fh(32), SLATE6)
    txt(s, fx(160), fy(482), fw(815), In(0.20),
        [[R("HARNESS 统一能力底座（场景 1 / 2 / 3 共享）", 8, AMBER, True)]], align=PP_ALIGN.CENTER)
    node(s, 190, 510, 176, 54, "上下文 · 记忆", "Context · Memory", F_VIOLET, VIOLET)
    node(s, 383, 510, 176, 54, "工具 / MCP 路由", "Tool Router", F_ORANGE, ORANGE)
    node(s, 576, 510, 176, 54, "知识 · SKILL 加载", "RAG · Skill Loader", F_CYAN, CYAN)
    node(s, 769, 510, 176, 54, "权限 · 安全护栏", "Guardrails · Audit", F_ROSE, ROSE)

    # ---------- Legend（边界之外） ----------
    txt(s, fx(180), fy(606), fw(120), In(0.18), [[R("Legend", 8, WHITE, True)]])
    legend = [(F_CYAN, CYAN, "交互 / 入口", 180), (F_EMERALD, EMERALD, "LLM / Agent", 300),
              (F_VIOLET, VIOLET, "知识 / 记忆", 420), (F_ORANGE, ORANGE, "协同 / 工具总线", 540),
              (F_ROSE, ROSE, "安全护栏", 690)]
    for (f, c, lab_t, lx) in legend:
        box(s, fx(lx), fy(630), In(0.17), In(0.10), f, c, 0.75)
        txt(s, fx(lx + 22), fy(626), fw(115), In(0.16), [[R(lab_t, 7, SLATE)]])
    box(s, fx(800), fy(630), In(0.17), In(0.10), None, AMBER, 0.75, dash="dash")
    txt(s, fx(822), fy(626), fw(160), In(0.16), [[R("Harness 边界（场景 4）", 7, SLATE)]])

    # ---------- 右侧摘要卡片 ----------
    CX, CW_ = In(9.78), In(13.333 - 9.78 - 0.35)
    card(s, CX, In(1.02), CW_, In(1.92), CYAN, "场景演进 1 → 2 → 3", [
        "场景 1：LLM + MCP 工具，简单问答与内容生成",
        "场景 2：LLM + RAG/SKILL，知识定边界、流程有章法",
        "场景 3：多 Agent 调度、协同、自适应",
        "由简到繁，能力逐级叠加而非互相替代"])
    card(s, CX, In(3.08), CW_, In(1.92), AMBER, "场景 4 · Harness 的内涵", [
        "一个大框把场景 1 / 2 / 3 全部包含在内",
        "统一上下文与记忆，跨场景共享状态",
        "统一工具 / MCP 路由与知识 · SKILL 加载",
        "统一多 Agent 调度与生命周期管理"])
    card(s, CX, In(5.14), CW_, In(1.92), ROSE, "运行与治理", [
        "权限最小化：工具与数据访问经护栏审批",
        "全链路审计：每次调用可观测、可回放",
        "自适应反馈：执行结果回流调度层动态调整",
        "结果回流业务：产出交付给企业用户与系统"])

    # ---------- 页脚 ----------
    txt(s, In(0.35), In(7.16), In(12.63), In(0.24),
        [[R("Agentic AI 原生架构 • 场景 1 工具 / 场景 2 知识 / 场景 3 协同 / 场景 4 Harness • chenliang · 2026-07",
            7, FAINT)]], align=PP_ALIGN.CENTER)

    save(prs, os.path.join(OUT, "Agentic原生架构图.pptx"))


if __name__ == "__main__":
    build()
