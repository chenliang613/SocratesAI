# -*- coding: utf-8 -*-
"""
生成《SaaS的智能化思考 —— 从 Salesforce Headless360 看软件的"去UI化"》PPT
"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

# ---------------- 设计系统 ----------------
INK      = RGBColor(0x0B, 0x1F, 0x33)   # 深蓝墨
NAVY     = RGBColor(0x10, 0x2A, 0x43)
BLUE     = RGBColor(0x1B, 0x6E, 0xF3)   # 主蓝（Salesforce 蓝调）
SKY      = RGBColor(0x35, 0xB8, 0xE0)
TEAL     = RGBColor(0x14, 0xB8, 0xA6)
AMBER    = RGBColor(0xF5, 0x9E, 0x0B)
CORAL    = RGBColor(0xF4, 0x6A, 0x5E)
PURPLE   = RGBColor(0x7C, 0x5C, 0xFF)
GREY     = RGBColor(0x5B, 0x6B, 0x7B)
LGREY    = RGBColor(0x8A, 0x98, 0xA6)
LINE     = RGBColor(0xD7, 0xDE, 0xE6)
PANEL    = RGBColor(0xF4, 0xF7, 0xFB)
PANEL2   = RGBColor(0xEC, 0xF1, 0xF8)
WHITE    = RGBColor(0xFF, 0xFF, 0xFF)

FONT = "Microsoft YaHei"

prs = Presentation()
prs.slide_width  = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]


def slide():
    return prs.slides.add_slide(BLANK)


def rect(s, x, y, w, h, fill, line=None, line_w=0.75, shadow=False, round_=False):
    shp = s.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if round_ else MSO_SHAPE.RECTANGLE,
        x, y, w, h)
    if round_:
        try:
            shp.adjustments[0] = 0.06
        except Exception:
            pass
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid(); shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line; shp.line.width = Pt(line_w)
    shp.shadow.inherit = False
    if shadow:
        el = shp._element.spPr
        ef = el.makeelement(qn('a:effectLst'), {})
        sh = ef.makeelement(qn('a:outerShdw'),
                            {'blurRad':'90000','dist':'40000','dir':'5400000','rotWithShape':'0'})
        clr = sh.makeelement(qn('a:srgbClr'), {'val':'0B1F33'})
        alpha = clr.makeelement(qn('a:alpha'), {'val':'18000'})
        clr.append(alpha); sh.append(clr); ef.append(sh); el.append(ef)
    return shp


def txt(s, x, y, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
        space_after=6, line_spacing=1.0, wrap=True):
    """runs: list of paragraphs; each paragraph = list of (text, size, color, bold, italic)"""
    tb = s.shapes.add_textbox(x, y, w, h); tf = tb.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    for m in (tf.margin_left, ):
        pass
    tf.margin_left = 0; tf.margin_right = 0; tf.margin_top = 0; tf.margin_bottom = 0
    for i, para in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align; p.space_after = Pt(space_after); p.space_before = Pt(0)
        p.line_spacing = line_spacing
        for (t, sz, c, b, it) in para:
            r = p.add_run(); r.text = t
            r.font.size = Pt(sz); r.font.color.rgb = c; r.font.bold = b
            r.font.italic = it; r.font.name = FONT
    return tb


def R(t, sz, c, b=False, it=False):
    return (t, sz, c, b, it)


def page_header(s, kicker, title, idx, total=15):
    rect(s, 0, 0, SW, Inches(1.18), WHITE)
    rect(s, Inches(0.6), Inches(0.34), Inches(0.12), Inches(0.5), BLUE)
    txt(s, Inches(0.85), Inches(0.26), Inches(10.5), Inches(0.32),
        [[R(kicker, 11.5, BLUE, True)]])
    txt(s, Inches(0.85), Inches(0.52), Inches(11.3), Inches(0.5),
        [[R(title, 23, INK, True)]])
    txt(s, SW - Inches(1.4), Inches(0.42), Inches(0.9), Inches(0.4),
        [[R(f"{idx:02d}", 17, BLUE, True), R(f" / {total}", 11, LGREY)]],
        align=PP_ALIGN.RIGHT)
    rect(s, 0, Inches(1.18), SW, Pt(1.2), LINE)
    rect(s, 0, SH - Inches(0.32), SW, Inches(0.32), INK)
    txt(s, Inches(0.6), SH - Inches(0.30), Inches(7), Inches(0.28),
        [[R("SaaS 的智能化思考  ·  软件去UI化深度分析", 8.5, RGBColor(0xC2,0xCE,0xDB))]],
        anchor=MSO_ANCHOR.MIDDLE)


def chip(s, x, y, w, label, color, h=Inches(0.34)):
    rect(s, x, y, w, h, color, round_=True)
    txt(s, x, y, w, h, [[R(label, 11, WHITE, True)]],
        align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


# ============================================================== #
# 1 封面
# ============================================================== #
s = slide()
rect(s, 0, 0, SW, SH, INK)
# 抽象网格/几何
rect(s, Inches(8.6), 0, Inches(4.73), SH, NAVY)
for i, (cx, cy, d, col) in enumerate([
    (Inches(9.4), Inches(1.2), Inches(1.5), BLUE),
    (Inches(11.3), Inches(2.6), Inches(2.4), RGBColor(0x18,0x3A,0x5C)),
    (Inches(9.9), Inches(4.4), Inches(1.9), TEAL),
    (Inches(11.8), Inches(5.4), Inches(1.2), AMBER)]):
    o = s.shapes.add_shape(MSO_SHAPE.OVAL, cx, cy, d, d)
    o.fill.solid(); o.fill.fore_color.rgb = col; o.line.fill.background()
    o.shadow.inherit = False
    if i in (1,):
        o.fill.background(); o.line.color.rgb = RGBColor(0x2A,0x4D,0x70); o.line.width=Pt(1.3)

rect(s, Inches(0.9), Inches(1.15), Inches(0.7), Pt(4), BLUE)
txt(s, Inches(0.9), Inches(1.35), Inches(7.5), Inches(0.5),
    [[R("DEEP ANALYSIS · SaaS × AI", 13, SKY, True)]])
txt(s, Inches(0.9), Inches(2.05), Inches(7.6), Inches(2.4),
    [[R("SaaS 的", 44, WHITE, True)],
     [R("智能化思考", 54, WHITE, True)]], line_spacing=1.0, space_after=2)
txt(s, Inches(0.9), Inches(4.45), Inches(7.4), Inches(1.0),
    [[R("从 Salesforce ", 19, RGBColor(0xC8,0xD6,0xE6)),
      R("Headless 360", 19, AMBER, True),
      R(" 看软件的「去 UI 化」", 19, RGBColor(0xC8,0xD6,0xE6))]])
rect(s, Inches(0.92), Inches(5.4), Inches(4.6), Pt(1.2), RGBColor(0x33,0x52,0x73))
txt(s, Inches(0.9), Inches(5.65), Inches(7.5), Inches(0.9),
    [[R("界面消失之处，智能开始生长 —— 当 Agent 成为新入口，", 12.5, RGBColor(0x9F,0xB2,0xC6))],
     [R("SaaS 的价值正从「画面」回归「能力」。", 12.5, RGBColor(0x9F,0xB2,0xC6))]],
    line_spacing=1.15)
txt(s, Inches(0.9), Inches(6.75), Inches(7), Inches(0.4),
    [[R("汇报人：chenliang   |   2026-06", 11, RGBColor(0x6E,0x84,0x9B))]])

# ============================================================== #
# 2 目录
# ============================================================== #
s = slide()
page_header(s, "CONTENTS", "目录 · 五大模块", 2)
items = [
    ("01", "趋势判断", "SaaS 演进到了「界面拐点」", BLUE),
    ("02", "概念解构", "Headless 360 与「去UI化」的本质", TEAL),
    ("03", "架构透视", "API-first / Composable / Agent-native", PURPLE),
    ("04", "商业冲击", "入口、护城河与定价的重构", AMBER),
    ("05", "落地路径", "企业的智能化转型四步走", CORAL),
]
x0 = Inches(0.85); y0 = Inches(1.65); cw = Inches(11.6); ch = Inches(0.92); gap = Inches(0.12)
for i,(no,t,d,c) in enumerate(items):
    y = y0 + i*(ch+gap)
    rect(s, x0, y, cw, ch, PANEL, round_=True)
    rect(s, x0, y, Inches(0.09), ch, c, round_=True)
    txt(s, x0+Inches(0.35), y, Inches(1.1), ch, [[R(no, 26, c, True)]],
        anchor=MSO_ANCHOR.MIDDLE)
    txt(s, x0+Inches(1.55), y+Inches(0.14), Inches(3.2), Inches(0.66),
        [[R(t, 16, INK, True)]], anchor=MSO_ANCHOR.MIDDLE)
    rect(s, x0+Inches(4.7), y+Inches(0.2), Pt(1.2), Inches(0.52), LINE)
    txt(s, x0+Inches(5.0), y, Inches(6.4), ch, [[R(d, 13.5, GREY)]],
        anchor=MSO_ANCHOR.MIDDLE)

# ============================================================== #
# 3 趋势：界面拐点
# ============================================================== #
s = slide()
page_header(s, "01 趋势判断", "SaaS 正在抵达「界面拐点」", 3)
txt(s, Inches(0.85), Inches(1.35), Inches(11.6), Inches(0.5),
    [[R("过去二十年 SaaS 的竞争核心是「更好的界面」；未来十年，竞争核心是「不需要界面」。",
        13.5, GREY)]])
# 三段时间轴
stages = [
    ("Web SaaS\n2005-2015", "把软件搬上云端\n以「页面+表单」承载业务，人适应软件", GREY, "人找功能"),
    ("Mobile / API SaaS\n2015-2023", "移动端 + 开放 API\n界面碎片化，集成成为刚需", BLUE, "功能进流程"),
    ("Agentic SaaS\n2023→", "AI Agent 成为操作主体\n意图驱动，UI 退居「可选层」", TEAL, "意图调用能力"),
]
x0 = Inches(0.85); cw = Inches(3.7); gap = Inches(0.28); y = Inches(2.1); ch = Inches(3.55)
rect(s, x0, y+Inches(1.62), Inches(11.6), Pt(3), LINE)
for i,(t,d,c,tag) in enumerate(stages):
    x = x0 + i*(cw+gap)
    o = s.shapes.add_shape(MSO_SHAPE.OVAL, x+cw/2-Inches(0.13), y+Inches(1.5), Inches(0.26), Inches(0.26))
    o.fill.solid(); o.fill.fore_color.rgb=c; o.line.color.rgb=WHITE; o.line.width=Pt(2.5); o.shadow.inherit=False
    rect(s, x, y, cw, Inches(1.35), WHITE, line=c, line_w=1.4, round_=True, shadow=True)
    txt(s, x+Inches(0.2), y+Inches(0.16), cw-Inches(0.4), Inches(1.1),
        [[R(t.split(chr(10))[0], 15, INK, True)],[R(t.split(chr(10))[1], 11, c, True)]],
        align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, space_after=2)
    rect(s, x, y+Inches(2.05), cw, Inches(1.5), PANEL, round_=True)
    txt(s, x+Inches(0.22), y+Inches(2.2), cw-Inches(0.44), Inches(1.2),
        [[R(d.split(chr(10))[0], 12.5, INK, True)],[R(d.split(chr(10))[1], 11, GREY)]],
        align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, space_after=4)
    chip(s, x+cw/2-Inches(0.85), y+Inches(3.62), Inches(1.7), tag, c)

txt(s, Inches(0.85), Inches(6.55), Inches(11.6), Inches(0.5),
    [[R("核心信号：", 13, CORAL, True),
      R("交互正从「人点击界面」迁移到「Agent 调用能力」，界面从必需品变成体验选项。",
        13, INK)]])

# ============================================================== #
# 4 什么是 Headless 360
# ============================================================== #
s = slide()
page_header(s, "02 概念解构", "什么是 Headless 360？", 4)
txt(s, Inches(0.85), Inches(1.32), Inches(11.6), Inches(0.7),
    [[R("Salesforce Customer 360 是「客户全景能力中台」；", 13.5, INK, True),
      R("Headless 则意味着把「头」（前端界面）从「身体」（业务能力）上解耦。", 13.5, GREY)]])
# 对照：传统 vs Headless
bx = Inches(0.85); bw = Inches(5.55); by = Inches(2.2); bh = Inches(3.0)
rect(s, bx, by, bw, bh, PANEL, line=LINE, round_=True)
rect(s, bx, by, bw, Inches(0.6), GREY, round_=True)
txt(s, bx, by, bw, Inches(0.6), [[R("传统一体化 SaaS", 14, WHITE, True)]],
    align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
for i,(t) in enumerate(["UI 与逻辑、数据强耦合","前端固定，体验由厂商定义",
                        "扩展靠插件，集成成本高","「软件 = 那个网页」"]):
    yy = by+Inches(0.78)+i*Inches(0.52)
    rect(s, bx+Inches(0.3), yy+Inches(0.07), Inches(0.12), Inches(0.12), GREY)
    txt(s, bx+Inches(0.55), yy, bw-Inches(0.8), Inches(0.4), [[R(t, 12.5, INK)]])

bx2 = Inches(6.9)
rect(s, bx2, by, bw, bh, PANEL2, line=TEAL, line_w=1.4, round_=True, shadow=True)
rect(s, bx2, by, bw, Inches(0.6), TEAL, round_=True)
txt(s, bx2, by, bw, Inches(0.6), [[R("Headless 360", 14, WHITE, True)]],
    align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
for i,(t) in enumerate(["能力以 API / 事件 / Agent 暴露","前端任意：网页、IM、语音、Agent",
                        "可组合（Composable）按需拼装","「软件 = 一组可被调用的能力」"]):
    yy = by+Inches(0.78)+i*Inches(0.52)
    rect(s, bx2+Inches(0.3), yy+Inches(0.07), Inches(0.12), Inches(0.12), TEAL)
    txt(s, bx2+Inches(0.55), yy, bw-Inches(0.8), Inches(0.4), [[R(t, 12.5, INK)]])

rect(s, Inches(0.85), Inches(5.5), Inches(11.6), Inches(1.05), INK, round_=True)
txt(s, Inches(1.2), Inches(5.5), Inches(11.0), Inches(1.05),
    [[R("一句话定义：", 14, AMBER, True),
      R("Headless 360 = 把「客户全景的业务能力」做成无界面的服务，", 14, WHITE),
      R("让任意终端、任意 AI Agent 都能即取即用 ——", 14, WHITE)],
     [R("界面不再是软件的入口，", 13, RGBColor(0xC8,0xD6,0xE6)),
      R("能力本身才是。", 13, SKY, True)]],
    anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.15, space_after=4)

# ============================================================== #
# 5 去UI化的本质
# ============================================================== #
s = slide()
page_header(s, "02 概念解构", "「去UI化」的本质：从人机界面到意图界面", 5)
txt(s, Inches(0.85), Inches(1.32), Inches(11.6), Inches(0.5),
    [[R("去UI化 ≠ 没有界面，而是界面从「主角」退为「可替换的呈现层」，交互的重心转移。", 13.5, GREY)]])

cards = [
    ("交互主体", "人 → Agent", "操作软件的主体从人，变为代表人的 AI 智能体", BLUE),
    ("交互方式", "点击 → 表达意图", "从「找按钮、填表单」到「说目标、给约束」", TEAL),
    ("界面角色", "入口 → 呈现层", "GUI 退化为可选的结果展示与确认环节", PURPLE),
    ("价值载体", "画面 → 能力", "竞争力从「界面好不好用」转向「能力强不强、可不可被调用」", AMBER),
]
x0=Inches(0.85); y0=Inches(2.05); cw=Inches(5.6); ch=Inches(1.78); gx=Inches(0.4); gy=Inches(0.28)
for i,(h,big,d,c) in enumerate(cards):
    x = x0 + (i%2)*(cw+gx); y = y0 + (i//2)*(ch+gy)
    rect(s, x, y, cw, ch, WHITE, line=LINE, round_=True, shadow=True)
    rect(s, x, y, Inches(0.1), ch, c, round_=True)
    txt(s, x+Inches(0.35), y+Inches(0.18), cw-Inches(0.6), Inches(0.35),
        [[R(h, 12, c, True)]])
    txt(s, x+Inches(0.35), y+Inches(0.5), cw-Inches(0.6), Inches(0.55),
        [[R(big, 21, INK, True)]])
    txt(s, x+Inches(0.35), y+Inches(1.12), cw-Inches(0.6), Inches(0.6),
        [[R(d, 12, GREY)]], line_spacing=1.1)

txt(s, Inches(0.85), Inches(6.55), Inches(11.6), Inches(0.5),
    [[R("洞察：", 13, CORAL, True),
      R("UI 是「人类带宽」的妥协；当 Agent 能以结构化方式高速调用能力，UI 的存在理由被部分抽空。",
        12.5, INK)]])

# ============================================================== #
# 6 驱动力
# ============================================================== #
s = slide()
page_header(s, "02 概念解构", "为什么「去UI化」此刻成为趋势？", 6)
drivers = [
    ("LLM/Agent 成熟", "自然语言成为通用接口，机器能理解意图、自主编排工具", BLUE),
    ("API 经济沉淀", "十年 SaaS 把业务能力 API 化，已具备「被调用」的底座", TEAL),
    ("MCP / 工具协议", "标准化的工具调用协议，让能力可被 Agent 自动发现与使用", PURPLE),
    ("企业降本诉求", "减少人工点击与重复操作，用自动化吃掉「界面劳动」", AMBER),
    ("多端碎片化", "网页/IM/语音/IoT 终端激增，统一前端不再现实", CORAL),
    ("数据飞轮", "调用即数据，能力被用得越多越聪明，正反馈强化", TEAL),
]
x0=Inches(0.85); y0=Inches(1.55); cw=Inches(3.72); ch=Inches(1.72); gx=Inches(0.22); gy=Inches(0.24)
for i,(t,d,c) in enumerate(drivers):
    x = x0 + (i%3)*(cw+gx); y = y0 + (i//3)*(ch+gy)
    rect(s, x, y, cw, ch, PANEL, round_=True)
    o = s.shapes.add_shape(MSO_SHAPE.OVAL, x+Inches(0.28), y+Inches(0.26), Inches(0.5), Inches(0.5))
    o.fill.solid(); o.fill.fore_color.rgb=c; o.line.fill.background(); o.shadow.inherit=False
    txt(s, x+Inches(0.28), y+Inches(0.26), Inches(0.5), Inches(0.5),
        [[R(str(i+1), 16, WHITE, True)]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    txt(s, x+Inches(0.95), y+Inches(0.28), cw-Inches(1.15), Inches(0.5),
        [[R(t, 14, INK, True)]], anchor=MSO_ANCHOR.MIDDLE)
    txt(s, x+Inches(0.3), y+Inches(0.92), cw-Inches(0.6), Inches(0.72),
        [[R(d, 11.5, GREY)]], line_spacing=1.12)
txt(s, Inches(0.85), Inches(6.62), Inches(11.6), Inches(0.4),
    [[R("供给（Agent 能力）× 需求（降本增效）× 协议（可调用标准）三者共振，趋势由「可能」变为「正在发生」。",
        12.5, INK, True)]], align=PP_ALIGN.CENTER)

# ============================================================== #
# 7 去UI化层次模型
# ============================================================== #
s = slide()
page_header(s, "03 架构透视", "去UI化的五级成熟度模型", 7)
txt(s, Inches(0.85), Inches(1.32), Inches(11.6), Inches(0.4),
    [[R("软件并非一夜「无界面」，而是沿成熟度逐级演进 —— 越往上，UI 占比越低、Agent 自主度越高。",
        13, GREY)]])
levels = [
    ("L0", "纯界面", "所有操作靠人点击 GUI", GREY, Inches(2.4)),
    ("L1", "API 可调", "能力开放 API，但仍以人为主", BLUE, Inches(3.0)),
    ("L2", "Copilot 辅助", "界面内嵌助手，建议+一键执行", PURPLE, Inches(3.6)),
    ("L3", "Agent 编排", "Agent 跨功能自动完成任务流", TEAL, Inches(4.2)),
    ("L4", "意图原生", "用户只表达目标，UI 仅用于确认", AMBER, Inches(4.8)),
]
x0=Inches(0.85); baseY=Inches(6.05); bw=Inches(2.18); gx=Inches(0.18)
for i,(lv,t,d,c,bh) in enumerate(levels):
    x = x0 + i*(bw+gx); y = baseY-bh
    rect(s, x, y, bw, bh, WHITE, line=c, line_w=1.3, round_=True, shadow=True)
    rect(s, x, y, bw, Inches(0.6), c, round_=True)
    txt(s, x, y, bw, Inches(0.6), [[R(lv, 17, WHITE, True)]],
        align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    txt(s, x+Inches(0.12), y+Inches(0.72), bw-Inches(0.24), Inches(0.5),
        [[R(t, 13.5, INK, True)]], align=PP_ALIGN.CENTER)
    txt(s, x+Inches(0.15), y+Inches(1.2), bw-Inches(0.3), bh-Inches(1.3),
        [[R(d, 11, GREY)]], align=PP_ALIGN.CENTER, line_spacing=1.12)
# 箭头
txt(s, Inches(0.85), Inches(6.25), Inches(11.6), Inches(0.35),
    [[R("UI 主导", 11, GREY, True),
      R("  ──────────────  界面占比下降 / 自主度上升  ──────────────►  ", 11, BLUE),
      R("能力主导", 11, AMBER, True)]], align=PP_ALIGN.CENTER)
txt(s, Inches(0.85), Inches(6.72), Inches(11.6), Inches(0.35),
    [[R("Headless 360 的定位：把企业能力推到 L3–L4，让 SaaS 成为「可被 Agent 调用的能力网络」。",
        12, INK, True)]], align=PP_ALIGN.CENTER)

# ============================================================== #
# 8 技术架构
# ============================================================== #
s = slide()
page_header(s, "03 架构透视", "Agent-native 的技术架构分层", 8)
layers = [
    ("交互层  Interaction", "网页 · 移动 · IM · 语音 · IoT · 第三方 Agent —— 多入口、可插拔", SKY),
    ("智能体层  Agent", "意图理解 · 任务规划 · 工具编排 · 记忆 · 护栏与审批", PURPLE),
    ("能力层  Capability API", "Headless 服务：订单/客户/营销/服务，以 API + 事件 + MCP 工具暴露", BLUE),
    ("数据层  Data 360", "统一客户数据 · 实时事件流 · 语义层 / 知识库（RAG）", TEAL),
    ("信任层  Trust & Governance", "权限 · 审计 · 可解释 · 合规 —— 贯穿全栈的横切关注点", AMBER),
]
x0=Inches(0.95); y0=Inches(1.55); w=Inches(11.4); h=Inches(0.86); gap=Inches(0.13)
for i,(t,d,c) in enumerate(layers):
    y = y0+i*(h+gap)
    if i==4:  # 信任层横切，画成右侧竖条概念——简化为整条强调
        rect(s, x0, y, w, h, INK, round_=True)
        rect(s, x0, y, Inches(0.12), h, c, round_=True)
        txt(s, x0+Inches(0.4), y, Inches(3.4), h, [[R(t, 14.5, c, True)]], anchor=MSO_ANCHOR.MIDDLE)
        txt(s, x0+Inches(3.9), y, w-Inches(4.2), h, [[R(d, 12, WHITE)]], anchor=MSO_ANCHOR.MIDDLE)
    else:
        rect(s, x0, y, w, h, PANEL, line=LINE, round_=True)
        rect(s, x0, y, Inches(0.12), h, c, round_=True)
        txt(s, x0+Inches(0.4), y, Inches(3.4), h, [[R(t, 14.5, INK, True)]], anchor=MSO_ANCHOR.MIDDLE)
        txt(s, x0+Inches(3.9), y, w-Inches(4.2), h, [[R(d, 12, GREY)]], anchor=MSO_ANCHOR.MIDDLE)
txt(s, Inches(0.95), Inches(6.95), Inches(11.4), Inches(0.4),
    [[R("关键转变：", 12.5, CORAL, True),
      R("应用不再是「页面的集合」，而是「能力 + 智能体 + 数据」的可组合系统，界面只是众多消费者之一。",
        12, INK)]])

# ============================================================== #
# 9 AI Agent 作为新UI
# ============================================================== #
s = slide()
page_header(s, "03 架构透视", "AI Agent：界面消失后的「新入口」", 9)
txt(s, Inches(0.85), Inches(1.32), Inches(11.6), Inches(0.45),
    [[R("当 UI 退场，谁来承接用户？答案是 Agent —— 它既是交互界面，也是任务执行者。", 13.5, GREY)]])
# 左：旧路径 右：新路径
lx=Inches(0.85); lw=Inches(5.5); ly=Inches(2.1); lh=Inches(3.1)
rect(s, lx, ly, lw, lh, PANEL, line=LINE, round_=True)
txt(s, lx, ly+Inches(0.18), lw, Inches(0.4), [[R("旧：人驱动 UI", 14, GREY, True)]], align=PP_ALIGN.CENTER)
for i,t in enumerate(["人产生需求","登录系统、找到模块","逐屏点击、填写表单","人工核对、提交","跨系统重复上述过程"]):
    yy=ly+Inches(0.72)+i*Inches(0.46)
    rect(s, lx+Inches(0.45), yy, lw-Inches(0.9), Inches(0.36), WHITE, line=LINE, round_=True)
    txt(s, lx+Inches(0.45), yy, lw-Inches(0.9), Inches(0.36), [[R(f"{i+1}. {t}", 11.5, INK)]],
        align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

rx=Inches(6.95);
rect(s, rx, ly, lw, lh, PANEL2, line=TEAL, line_w=1.4, round_=True, shadow=True)
txt(s, rx, ly+Inches(0.18), lw, Inches(0.4), [[R("新：意图驱动 Agent", 14, TEAL, True)]], align=PP_ALIGN.CENTER)
for i,t in enumerate(["人表达目标与约束","Agent 理解意图、拆解任务","自动调用 Headless 能力","跨系统自主编排、回填","只在关键节点请人确认"]):
    yy=ly+Inches(0.72)+i*Inches(0.46)
    rect(s, rx+Inches(0.45), yy, lw-Inches(0.9), Inches(0.36), WHITE, line=TEAL, round_=True)
    txt(s, rx+Inches(0.45), yy, lw-Inches(0.9), Inches(0.36), [[R(f"{i+1}. {t}", 11.5, INK)]],
        align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

txt(s, Inches(0.85), Inches(6.6), Inches(11.6), Inches(0.5),
    [[R("结果：", 13, CORAL, True),
      R("交互步骤从「N 屏点击」压缩为「1 句意图 + 1 次确认」，软件价值从「好操作」变为「会办事」。",
        12.5, INK)]])

# ============================================================== #
# 10 商业冲击
# ============================================================== #
s = slide()
page_header(s, "04 商业冲击", "去UI化如何重构 SaaS 商业模式", 10)
cols = [
    ("入口之争", "界面入口 → Agent 入口", "谁掌握 Agent，谁掌握用户；SaaS 可能沦为「被调用的后端」", CORAL),
    ("护城河", "UI/体验壁垒 → 能力/数据壁垒", "界面易被复刻，独有数据与可靠能力才是新护城河", BLUE),
    ("定价模型", "按席位 → 按调用/按结果", "Agent 自动执行下，按用量、按完成的业务结果计费成为主流", PURPLE),
    ("产品形态", "应用 App → 能力 + 工具集", "卖的不再是「界面应用」，而是可被编排的能力与工具", TEAL),
]
x0=Inches(0.85); y0=Inches(1.55); cw=Inches(5.6); ch=Inches(2.35); gx=Inches(0.4); gy=Inches(0.28)
for i,(h,shift,d,c) in enumerate(cols):
    x=x0+(i%2)*(cw+gx); y=y0+(i//2)*(ch+gy)
    rect(s, x, y, cw, ch, WHITE, line=LINE, round_=True, shadow=True)
    rect(s, x, y, cw, Inches(0.55), c, round_=True)
    txt(s, x, y, cw, Inches(0.55), [[R(h, 14.5, WHITE, True)]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    txt(s, x+Inches(0.35), y+Inches(0.72), cw-Inches(0.7), Inches(0.5),
        [[R(shift, 15, INK, True)]])
    rect(s, x+Inches(0.35), y+Inches(1.28), cw-Inches(0.7), Pt(1), LINE)
    txt(s, x+Inches(0.35), y+Inches(1.42), cw-Inches(0.7), Inches(0.85),
        [[R(d, 12.5, GREY)]], line_spacing=1.18)
txt(s, Inches(0.85), Inches(6.72), Inches(11.6), Inches(0.4),
    [[R("最大风险：", 12.5, CORAL, True),
      R("被「平台级 Agent」中介化，失去与终端用户的直接触点 —— 必须主动成为 Agent 生态的能力供给方。",
        12, INK)]])

# ============================================================== #
# 11 机会与风险
# ============================================================== #
s = slide()
page_header(s, "04 商业冲击", "机会与风险：一体两面", 11)
ox=Inches(0.85); ow=Inches(5.6); oy=Inches(1.55); oh=Inches(4.55)
rect(s, ox, oy, ow, oh, RGBColor(0xEC,0xFB,0xF6), line=TEAL, line_w=1.3, round_=True)
rect(s, ox, oy, ow, Inches(0.6), TEAL, round_=True)
txt(s, ox, oy, ow, Inches(0.6), [[R("机会  Opportunity", 15, WHITE, True)]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
for i,t in enumerate([
    "把能力 API 化，进入每一个 Agent 工作流",
    "按结果计费，打开更高价值的定价空间",
    "数据飞轮：被调用越多，能力越强、越独占",
    "极致降本：吃掉企业内部的「界面劳动」",
    "多端无界：一套能力服务所有终端与场景"]):
    yy=oy+Inches(0.8)+i*Inches(0.72)
    rect(s, ox+Inches(0.35), yy+Inches(0.06), Inches(0.16), Inches(0.16), TEAL)
    txt(s, ox+Inches(0.65), yy, ow-Inches(1.0), Inches(0.62), [[R(t, 12.5, INK)]], line_spacing=1.1)

rx=Inches(6.95)
rect(s, rx, oy, ow, oh, RGBColor(0xFD,0xF0,0xEE), line=CORAL, line_w=1.3, round_=True)
rect(s, rx, oy, ow, Inches(0.6), CORAL, round_=True)
txt(s, rx, oy, ow, Inches(0.6), [[R("风险  Risk", 15, WHITE, True)]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
for i,t in enumerate([
    "被平台 Agent 中介化，丢失用户直接触点",
    "界面消失 → 品牌与体验差异被削弱",
    "Agent 自主执行带来的可控性与合规风险",
    "幻觉 / 误操作，对关键业务的信任门槛高",
    "能力被比价：沦为可替换的「调用后端」"]):
    yy=oy+Inches(0.8)+i*Inches(0.72)
    rect(s, rx+Inches(0.35), yy+Inches(0.06), Inches(0.16), Inches(0.16), CORAL)
    txt(s, rx+Inches(0.65), yy, ow-Inches(1.0), Inches(0.62), [[R(t, 12.5, INK)]], line_spacing=1.1)

# ============================================================== #
# 12 落地路径
# ============================================================== #
s = slide()
page_header(s, "05 落地路径", "企业智能化转型：四步走", 12)
steps = [
    ("STEP 1", "能力 API 化", "梳理核心业务能力，统一为 API / 事件 / 工具，先「可被调用」", BLUE),
    ("STEP 2", "数据 360 化", "打通数据孤岛，建语义层与知识库，喂养可靠的 Agent", TEAL),
    ("STEP 3", "Copilot 嵌入", "在现有界面内嵌助手，低风险验证意图驱动的价值", PURPLE),
    ("STEP 4", "Agent 编排", "上线跨系统自主 Agent，配护栏与审批，迈向 L3–L4", AMBER),
]
x0=Inches(0.85); y0=Inches(1.65); cw=Inches(2.78); ch=Inches(3.7); gx=Inches(0.18)
for i,(st,t,d,c) in enumerate(steps):
    x=x0+i*(cw+gx)
    rect(s, x, y0, cw, ch, WHITE, line=LINE, round_=True, shadow=True)
    rect(s, x, y0, cw, Inches(1.0), c, round_=True)
    txt(s, x, y0+Inches(0.14), cw, Inches(0.35), [[R(st, 12, WHITE, True)]], align=PP_ALIGN.CENTER)
    txt(s, x, y0+Inches(0.46), cw, Inches(0.5), [[R(t, 16, WHITE, True)]], align=PP_ALIGN.CENTER)
    txt(s, x+Inches(0.25), y0+Inches(1.25), cw-Inches(0.5), Inches(2.3),
        [[R(d, 12.5, GREY)]], line_spacing=1.25)
    o=s.shapes.add_shape(MSO_SHAPE.OVAL, x+cw/2-Inches(0.35), y0+ch-Inches(0.9), Inches(0.7), Inches(0.7))
    o.fill.solid(); o.fill.fore_color.rgb=PANEL; o.line.color.rgb=c; o.line.width=Pt(1.5); o.shadow.inherit=False
    txt(s, x+cw/2-Inches(0.35), y0+ch-Inches(0.9), Inches(0.7), Inches(0.7),
        [[R(str(i+1), 22, c, True)]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    if i<3:
        txt(s, x+cw+Inches(0.0), y0+ch/2-Inches(0.2), Inches(0.18), Inches(0.4),
            [[R("›", 22, LGREY, True)]], align=PP_ALIGN.CENTER)
txt(s, Inches(0.85), Inches(6.7), Inches(11.6), Inches(0.4),
    [[R("原则：", 12.5, CORAL, True),
      R("从「内嵌助手」起步控制风险，以「能力+数据」为根基，逐级提升 Agent 自主度，不可一步到位。",
        12, INK)]])

# ============================================================== #
# 13 关键洞察总结
# ============================================================== #
s = slide()
page_header(s, "STRATEGIC TAKEAWAYS", "五个必须记住的判断", 13)
points = [
    ("UI 是带宽妥协", "界面是为「人类有限带宽」设计的；Agent 不需要它，去UI化是结构性必然。"),
    ("软件 = 能力网络", "未来的 SaaS 是「可被调用的能力 + 数据 + 智能体」，而非一堆页面。"),
    ("入口在迁移", "用户入口正从界面转向 Agent，谁不进 Agent 工作流，谁就失去用户。"),
    ("数据是新护城河", "界面可复刻、能力会被比价，独有数据与可靠执行才是壁垒。"),
    ("从结果计费", "Agent 自动办事，商业模式从「卖席位」转向「卖结果」。"),
]
x0=Inches(0.85); y0=Inches(1.6); w=Inches(11.6); h=Inches(0.92); gap=Inches(0.12)
cols_c=[BLUE,TEAL,PURPLE,AMBER,CORAL]
for i,(t,d) in enumerate(points):
    y=y0+i*(h+gap)
    rect(s, x0, y, w, h, PANEL, round_=True)
    o=s.shapes.add_shape(MSO_SHAPE.OVAL, x0+Inches(0.28), y+Inches(0.21), Inches(0.5), Inches(0.5))
    o.fill.solid(); o.fill.fore_color.rgb=cols_c[i]; o.line.fill.background(); o.shadow.inherit=False
    txt(s, x0+Inches(0.28), y+Inches(0.21), Inches(0.5), Inches(0.5),
        [[R(str(i+1), 17, WHITE, True)]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    txt(s, x0+Inches(1.0), y+Inches(0.12), Inches(3.0), h-Inches(0.24),
        [[R(t, 15, INK, True)]], anchor=MSO_ANCHOR.MIDDLE)
    rect(s, x0+Inches(3.95), y+Inches(0.2), Pt(1.2), h-Inches(0.4), LINE)
    txt(s, x0+Inches(4.2), y, w-Inches(4.5), h, [[R(d, 12.5, GREY)]],
        anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.1)

# ============================================================== #
# 14 未来展望
# ============================================================== #
s = slide()
rect(s, 0, 0, SW, SH, INK)
rect(s, 0, 0, SW, Inches(1.18), NAVY)
rect(s, Inches(0.6), Inches(0.34), Inches(0.12), Inches(0.5), AMBER)
txt(s, Inches(0.85), Inches(0.3), Inches(11), Inches(0.7),
    [[R("未来展望 · Beyond the Interface", 22, WHITE, True)]], anchor=MSO_ANCHOR.MIDDLE)
fut = [
    ("界面成为可选项", "GUI 不会消失，但从「唯一入口」降级为「可选的呈现与确认层」。", SKY),
    ("Agent 即操作系统", "企业内部出现「Agent 编排层」，统一调度各 SaaS 能力。", TEAL),
    ("能力市场化", "可被调用的能力像 API 一样在 Agent 生态中流通、被比价、被组合。", AMBER),
    ("人退到监督位", "人从「操作者」变为「目标设定者 + 关键决策的审批者」。", PURPLE),
]
x0=Inches(0.85); y0=Inches(1.65); cw=Inches(5.6); ch=Inches(2.2); gx=Inches(0.4); gy=Inches(0.3)
for i,(t,d,c) in enumerate(fut):
    x=x0+(i%2)*(cw+gx); y=y0+(i//2)*(ch+gy)
    rect(s, x, y, cw, ch, NAVY, line=RGBColor(0x2A,0x4D,0x70), round_=True)
    rect(s, x, y, Inches(0.1), ch, c, round_=True)
    txt(s, x+Inches(0.4), y+Inches(0.3), cw-Inches(0.7), Inches(0.5),
        [[R(t, 17, c, True)]])
    txt(s, x+Inches(0.4), y+Inches(0.95), cw-Inches(0.7), Inches(1.1),
        [[R(d, 13, RGBColor(0xC8,0xD6,0xE6))]], line_spacing=1.25)
txt(s, Inches(0.85), Inches(6.75), Inches(11.6), Inches(0.5),
    [[R("Headless 360 不是终点，而是 SaaS 走向「能力即服务、智能即交互」时代的起点。",
        13.5, AMBER, True)]], align=PP_ALIGN.CENTER)

# ============================================================== #
# 15 结语
# ============================================================== #
s = slide()
rect(s, 0, 0, SW, SH, INK)
rect(s, 0, Inches(2.7), SW, Inches(2.1), NAVY)
txt(s, Inches(1), Inches(1.4), Inches(11.3), Inches(0.5),
    [[R("CONCLUSION", 14, SKY, True)]], align=PP_ALIGN.CENTER)
txt(s, Inches(1), Inches(2.85), Inches(11.3), Inches(0.9),
    [[R("界面消失之处，智能开始生长。", 34, WHITE, True)]], align=PP_ALIGN.CENTER)
txt(s, Inches(1.5), Inches(3.85), Inches(10.3), Inches(0.8),
    [[R("当软件「去UI化」，竞争的胜负手不再是画面，", 15, RGBColor(0xC8,0xD6,0xE6))],
     [R("而是 —— 你的能力，是否值得被每一个 Agent 调用。", 15, RGBColor(0xC8,0xD6,0xE6))]],
    align=PP_ALIGN.CENTER, line_spacing=1.3)
rect(s, Inches(5.66), Inches(5.25), Inches(2), Pt(2), AMBER)
txt(s, Inches(1), Inches(5.5), Inches(11.3), Inches(0.5),
    [[R("SaaS 的智能化思考", 14, WHITE, True)]], align=PP_ALIGN.CENTER)
txt(s, Inches(1), Inches(6.05), Inches(11.3), Inches(0.4),
    [[R("Thanks  ·  欢迎讨论与指正   |   chenliang  ·  2026-06", 11.5, RGBColor(0x7E,0x93,0xAB))]],
    align=PP_ALIGN.CENTER)

out = "/Users/apple/SocratesAI/SaaS的智能化思考.pptx"
prs.save(out)
print("saved:", out, "slides:", len(prs.slides._sldIdLst))
